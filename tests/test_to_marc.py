import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from pymarc import MARCReader

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pipeline import ebsconet_to_marc as m  # noqa: E402
from pipeline.ebsconet_prep import load_config  # noqa: E402

COLS = ["Title Name", "ISSN", "Start Date", "Expiration Date", "Order Number",
        "Total Cost", "Publisher Name", "URL", "Account Number", "Format",
        "FOLIO Fund", "FOLIO Expense Class", "FOLIO Org", "Title Number"]
TAGS = ["245$a", "020$a", "990$s", "990$t", "990$o", "990$c", "264$a", "856$u",
        "990$a", "Ignore"]


@pytest.fixture
def cfg():
    return load_config(ROOT / "ebsconet_config.json")


@pytest.fixture
def headers(tmp_path):
    path = tmp_path / "hdr.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(TAGS)
    ws.append(COLS[:len(TAGS)])
    wb.save(path)
    return path


def row(**kw):
    base = {"Title Name": "Café Journal", "ISSN": "1234-5678",
            "Start Date": "2026-01-01", "Expiration Date": "2026-12-31",
            "Order Number": "U1", "Total Cost": 1234.5, "Publisher Name": "Pub $ Co",
            "URL": "http://x.org", "Account Number": "BR1", "Format": "Online Only",
            "FOLIO Fund": "F1", "FOLIO Expense Class": "E1", "FOLIO Org": "O1",
            "Title Number": "P1"}
    base.update(kw)
    return base


def test_parse_target():
    assert m.parse_target("990$s") == ("990", "s")
    assert m.parse_target("Ignore") is None
    assert m.parse_target(None) is None
    assert m.parse_target("bad") is None


def test_load_tag_map_skips_ignored_and_adds_extras(headers, cfg):
    tm = m.load_tag_map(headers, cfg)
    cols = {c for c, _, _ in tm}
    assert "Format" not in cols
    assert ("Title Name", "245", "a") in tm
    assert ("FOLIO Fund", "990", "f") in tm
    assert ("Title Number", "990", "i") in tm
    assert ("Title Number Type", "990", "j") in tm


def test_build_record_fields(headers, cfg):
    tm = m.load_tag_map(headers, cfg)
    rec = m.build_record(row(), tm, "U1", cfg)
    assert rec["001"].data == "U1"
    assert rec["245"]["a"] == "Café Journal"
    assert tuple(rec["245"].indicators) == ("0", "0")
    assert tuple(rec["264"].indicators) == (" ", "1")
    assert rec["020"]["a"] == "1234-5678"
    f990 = {s.code: s.value for s in rec["990"].subfields}
    assert f990 == {"s": "2026-01-01", "t": "2026-12-31", "o": "U1", "c": "1234.50",
                    "a": "BR1", "f": "F1", "e": "E1", "v": "O1", "i": "P1"}
    assert len(rec.get_fields("990")) == 1


def test_blank_values_omitted(headers, cfg):
    tm = m.load_tag_map(headers, cfg)
    rec = m.build_record(row(ISSN="", URL=None), tm, "U1", cfg)
    assert rec.get("020") is None and rec.get("856") is None


def test_control_ids_suffix_repeats():
    rows = [{"Order Number": "A"}, {"Order Number": "A"}, {"Order Number": "B"},
            {"Order Number": None}]
    assert m.control_ids(rows, "Order Number") == ["A", "A-2", "B", "ebsconet-row4"]


def test_mrk_escapes_dollar(headers, cfg):
    tm = m.load_tag_map(headers, cfg)
    text = m.to_mrk(m.build_record(row(), tm, "U1", cfg))
    assert "{dollar}" in text and "=245  00$aCafé Journal" in text
    assert "=264  \\1$aPub {dollar} Co" in text


def test_convert_all_round_trip(headers, cfg, tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(COLS)
    for r in (row(), row(**{"Title Name": "Two", "ISSN": None})):
        ws.append([r[c] for c in COLS])
    wb.save(src / cfg["output_names"]["online"])

    results = m.convert_all(src, cfg, headers)
    assert len(results) == 1
    assert results[0]["records"] == 2 and results[0]["problems"] == []
    mrc = src / "marc" / "library-EBSCONET_online.mrc"
    with open(mrc, "rb") as fh:
        recs = list(MARCReader(fh, to_unicode=True, force_utf8=True))
    assert [r["245"]["a"] for r in recs] == ["Café Journal", "Two"]
    assert recs[1].get("020") is None
    assert (src / "marc" / "library-EBSCONET_online.mrk").exists()


def test_empty_workbook_writes_no_files(headers, cfg, tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    wb = Workbook()
    wb.active.append(COLS)
    wb.save(src / cfg["output_names"]["print"])
    stale = src / "marc"
    stale.mkdir()
    (stale / "library-EBSCONET-print.mrc").write_bytes(b"old")
    results = m.convert_all(src, cfg, headers)
    assert results[0]["records"] == 0
    assert not (stale / "library-EBSCONET-print.mrc").exists()
