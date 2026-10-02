import re
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pipeline import ebsconet_prep as prep  # noqa: E402

HEADERS = ["Title Name", "ISSN", "Format", "Start Date", "Expiration Date",
           "Order Number", "Total Cost", "Publisher Package", "Publisher Name",
           "Subject Category", "Publisher Product Code", "Invoice Date"]


@pytest.fixture
def cfg():
    return prep.load_config(ROOT / "ebsconet_config.json")


def make_row(**kw):
    base = {"Title Name": "T", "ISSN": "1234-5678", "Format": "Online Only",
            "Start Date": "01/02/2026", "Expiration Date": "12/31/2026",
            "Order Number": "U1", "Total Cost": 10, "Publisher Package": "",
            "Publisher Name": "Pub", "Subject Category": "Physical Sciences",
            "Publisher Product Code": "P1", "Invoice Date": "", "_row": 2}
    base.update(kw)
    return base


def test_normalize_url():
    upper = "HTTP://WWW.JNCI.OUPJOURNALS.ORG"
    assert prep.normalize_url(upper) == "http://www.jnci.oupjournals.org"
    mixed = "HTTPS://Example.ORG/Path/CaseKept?Q=A"
    assert prep.normalize_url(mixed) == "https://example.org/Path/CaseKept?Q=A"
    assert prep.normalize_url(" www.example.org ") == "www.example.org"
    assert prep.normalize_url(None) == "" and prep.normalize_url("") == ""


def test_enrich_normalizes_url(cfg):
    out, _ = prep.enrich(make_row(URL="HTTP://WWW.X.ORG"), "online", cfg)
    assert out["URL"] == "http://www.x.org"


def test_cancellation_restriction_is_inverted():
    assert prep.cancellation_restriction("Yes") == "false"
    assert prep.cancellation_restriction(" no ") == "true"
    assert prep.cancellation_restriction(None) == ""
    assert prep.cancellation_restriction("maybe") == ""


def test_line_description():
    both = prep.line_description("Site License", "Monthly (8-14 issues)", "; ")
    assert both == "Site License; Monthly (8-14 issues)"
    assert prep.line_description(None, "Monthly", "; ") == "Monthly"
    assert prep.line_description("Single Site", "Not Applicable", "; ") == "Single Site"
    assert prep.line_description("", "not applicable", "; ") == ""


def test_enrich_adds_description_and_restriction(cfg):
    out, _ = prep.enrich(make_row(Cancellable="No", Descriptor="Single Site",
                                  Frequency="Quarterly (4 issues)"), "online", cfg)
    assert out["Cancellation Restriction"] == "true"
    assert out["PO Line Description"] == "Single Site; Quarterly (4 issues)"


def test_to_iso():
    assert prep.to_iso("01/02/2026") == "2026-01-02"
    assert prep.to_iso("2026-01-02") == "2026-01-02"
    assert prep.to_iso("") == ""
    assert prep.to_iso("garbage") is None


@pytest.mark.parametrize("fmt,route", [
    ("Online Only", "online"), ("Database", "online"), ("E-Books", "online"),
    ("Print", "print"), ("Online + Print", "pe"),
    ("Print + Online + Email", "pe"), ("Fee", "excluded"), ("Weird", None)])
def test_route_for(cfg, fmt, route):
    assert prep.route_for(fmt, cfg) == route


def test_classify_loadable(cfg):
    assert prep.classify(make_row(), cfg) == ("online", "")


def test_classify_zero_cost_package_member(cfg):
    row = make_row(**{"Total Cost": 0, "Publisher Package": "Big Deal"})
    assert prep.classify(row, cfg) == (None, "Package member at zero cost")


def test_usage_loading_service_lines_load_by_default(cfg):
    row = make_row(**{"Title Name": "Usage Loading Service - Ovid", "Total Cost": 125})
    assert prep.classify(row, cfg) == ("online", "")


def test_usage_loading_service_lines_can_be_excluded(cfg):
    cfg["rules"]["exclude_usage_loading_service"] = True
    row = make_row(**{"Title Name": "usage loading service - Ovid", "Total Cost": 125})
    assert prep.classify(row, cfg) == (None, "Usage Loading Service")
    other = make_row(**{"Title Name": "Journal of Usage Loading Services Research"})
    assert prep.classify(other, cfg) == ("online", "")      # only titles that start with it


def test_classify_zero_cost_excluded(cfg):
    assert prep.classify(make_row(**{"Total Cost": 0}), cfg) == (None, "Zero cost")
    assert prep.classify(make_row(Format="Print", **{"Total Cost": 0}), cfg)[1] == "Zero cost"


def test_classify_zero_cost_kept_when_rule_off(cfg):
    cfg["rules"]["exclude_zero_cost"] = False
    assert prep.classify(make_row(**{"Total Cost": 0}), cfg)[0] == "online"


def test_is_package_only_when_title_says_or_matches(cfg):
    assert prep.is_package(make_row(**{"Title Name": "Big Package"}), cfg)
    assert prep.is_package(make_row(**{"Title Name": "Current Protocols Collection"}), cfg)
    assert prep.is_package(make_row(**{"Title Name": "AACR Journals Suite"}), cfg)
    assert not prep.is_package(make_row(**{"Title Name": "Suitelife Letters"}), cfg)
    assert prep.is_package(make_row(**{"Title Name": "Wiley Core Collection",
                                       "Publisher Package": "wiley core collection"}), cfg)
    member = make_row(**{"Title Name": "Chem", "Publisher Package": "Cell Press"})
    assert not prep.is_package(member, cfg)
    assert prep.is_package_member(member, cfg)


def test_no_issn_rows_are_loaded_by_default(cfg):
    assert prep.classify(make_row(ISSN=""), cfg) == ("online", "")
    row = make_row(ISSN="", **{"Title Name": "Chem", "Publisher Package": "Cell Press"})
    assert prep.classify(row, cfg) == ("online", "")


def test_skip_missing_issn_rule_still_works_when_switched_on(cfg):
    cfg["rules"]["skip_missing_issn"] = True
    assert prep.classify(make_row(ISSN=""), cfg) == (None, "Missing ISSN")
    row = make_row(ISSN="", **{"Title Name": "Chem", "Publisher Package": "Cell Press"})
    assert prep.classify(row, cfg) == (None, "Missing ISSN")


def test_generated_identifier_only_when_no_issn_and_no_title_number(cfg):
    plain, _ = prep.enrich(make_row(), "online", cfg)
    assert plain["Generated ID?"] == "No" and plain["Title Number"] == "P1"
    no_issn, _ = prep.enrich(make_row(ISSN=""), "online", cfg)
    assert no_issn["Title Number"] == "P1" and no_issn["Generated ID?"] == "No"
    neither, _ = prep.enrich(make_row(ISSN="", **{"Publisher Product Code": ""}),
                             "online", cfg)
    assert neither["Title Number"] == "NOISSN-U1"
    assert neither["Title Number Type"] == "Local identifier"
    assert neither["Generated ID?"] == "Yes"
    has_issn, _ = prep.enrich(make_row(**{"Publisher Product Code": ""}), "online", cfg)
    assert has_issn["Title Number"] == "" and has_issn["Generated ID?"] == "No"


def test_classify_costed_package_exempt_from_issn_skip(cfg):
    cfg["rules"]["skip_missing_issn"] = True
    row = make_row(ISSN="", **{"Title Name": "AAAS Journals Package",
                               "Publisher Package": "AAAS Journals Package"})
    assert prep.classify(row, cfg)[0] == "online"


def test_classify_fee_and_unknown_format(cfg):
    assert prep.classify(make_row(Format="Fee"), cfg)[0] is None
    assert "Unrecognized" in prep.classify(make_row(Format="Zzz"), cfg)[1]


def test_enrich_adds_columns_and_iso_dates(cfg):
    out, warn = prep.enrich(make_row(), "print", cfg)
    assert out["Start Date"] == "2026-01-02"
    assert out["FOLIO Fund"] == "TEST-PRINT"
    assert out["FOLIO Expense Class"] == "PHY"
    assert out["FOLIO Access Provider"] == "EBSCO"
    assert out["Package?"] == "No"
    assert out["Title Number"] == "P1"
    assert out["Title Number Type"] == "Publisher or distributor number"
    blank_tn, _ = prep.enrich(make_row(**{"Publisher Product Code": ""}), "print", cfg)
    assert blank_tn["Title Number"] == "" and blank_tn["Title Number Type"] == ""
    assert warn == []


def test_expense_classes_can_be_switched_off(cfg):
    cfg["rules"]["use_expense_classes"] = False
    out, _ = prep.enrich(make_row(), "print", cfg)
    assert out["FOLIO Expense Class"] == ""


def test_enrich_bad_date_warns(cfg):
    out, warn = prep.enrich(make_row(**{"Start Date": "nope"}), "online", cfg)
    assert out["Start Date"] == "" and len(warn) == 1


def test_prepare_end_to_end(cfg, tmp_path):
    rows = [
        make_row(**{"Order Number": "A"}),
        make_row(**{"Order Number": "A", "Title Name": "T2", "ISSN": "1111-2222"}),
        make_row(Format="Print", **{"Order Number": "B"}),
        make_row(Format="Online + Print", **{"Order Number": "C"}),
        make_row(Format="Fee"),
        make_row(ISSN="", **{"Order Number": "D"}),
        make_row(**{"Total Cost": 0, "Publisher Package": "X"}),
    ]
    src = tmp_path / "sop.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r[h] for h in HEADERS])
    wb.save(src)

    hdr = tmp_path / "hdr.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["245$a"] + ["Ignore"] * (len(HEADERS) - 1))
    ws.append(HEADERS)
    wb.save(hdr)

    summary = prep.prepare(src, tmp_path / "out", cfg, hdr)
    assert summary["read"] == 7
    assert summary["routed"] == {"online": 3, "print": 1, "pe": 1}
    assert summary["excluded"] == 2            # the Fee row and the zero-cost package member
    assert summary["duplicate_orders"] == ["A"]

    no_issn = (tmp_path / "out" / "prep_no_issn.csv").read_text(encoding="utf-8")
    assert "D" in no_issn and no_issn.count("\n") == 2       # header + the one no-ISSN row
    report = (tmp_path / "out" / "prep_report.txt").read_text(encoding="utf-8")
    assert "WITHOUT an ISSN" in report

    out = load_workbook(tmp_path / "out" / cfg["output_names"]["online"])
    ws = out.active
    assert ws.max_row == 4
    names = [c.value for c in ws[1]]
    col = names.index("FOLIO Fund") + 1                        # an added column
    assert ws.cell(row=1, column=col).fill.start_color.rgb.endswith("FFFFCC")
    assert ws.cell(row=3, column=col).fill.start_color.rgb.endswith("FFFFCC")
    assert ws["A1"].fill.fill_type is None                     # SOP columns: no fill
    assert ws["A3"].fill.fill_type is None
    assert (tmp_path / "out" / "prep_report.txt").exists()
    assert (tmp_path / "out" / "prep_exclusions.csv").exists()


def test_config_is_valid_json():
    json.loads((ROOT / "ebsconet_config.json").read_text(encoding="utf-8"))


def test_customer_values_win_over_the_config_defaults(cfg):
    row = make_row(**{"FOLIO Fund": " MYFUND ", "FOLIO Expense Class": "MYEC",
                      "FOLIO Access Provider": "MYORG"})
    out, _ = prep.enrich(row, "online", cfg)
    assert (out["FOLIO Fund"], out["FOLIO Expense Class"], out["FOLIO Access Provider"]) == (
        "MYFUND", "MYEC", "MYORG")


def test_blank_or_missing_customer_values_fall_back_to_defaults(cfg):
    blank_row = make_row(**{"FOLIO Fund": "", "FOLIO Expense Class": None,
                            "FOLIO Access Provider": "  "})
    out, _ = prep.enrich(blank_row, "print", cfg)
    assert (out["FOLIO Fund"], out["FOLIO Expense Class"], out["FOLIO Access Provider"]) == (
        "TEST-PRINT", "PHY", "EBSCO")                       # subject map, default org
    out, _ = prep.enrich(make_row(), "online", cfg)         # no such columns at all
    assert out["FOLIO Fund"] == "TEST-ELEC"


def test_expense_class_off_ignores_the_customers_value_too(cfg):
    cfg["rules"]["use_expense_classes"] = False
    out, _ = prep.enrich(make_row(**{"FOLIO Expense Class": "MYEC"}), "online", cfg)
    assert out["FOLIO Expense Class"] == ""


def test_prepare_uses_existing_customer_columns_once_and_reports_defaults(cfg, tmp_path):
    headers = HEADERS + ["FOLIO Access Provider", "FOLIO Fund", "FOLIO Expense Class"]
    rows = [make_row(**{"Order Number": "A", "FOLIO Access Provider": "O1", "FOLIO Fund": "F1",
                        "FOLIO Expense Class": "E1"}),
            make_row(**{"Order Number": "B", "FOLIO Access Provider": "", "FOLIO Fund": "F2",
                        "FOLIO Expense Class": ""})]
    src = tmp_path / "sop.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    wb.save(src)
    prep.prepare(src, tmp_path / "out", cfg, None)

    out = load_workbook(tmp_path / "out" / cfg["output_names"]["online"]).active
    names = [c.value for c in out[1]]
    assert names.count("FOLIO Fund") == 1 and names.count("FOLIO Access Provider") == 1
    got = [dict(zip(names, [c.value for c in row])) for row in out.iter_rows(min_row=2)]
    keys = ("FOLIO Fund", "FOLIO Access Provider", "FOLIO Expense Class")
    assert [tuple(g[k] for k in keys) for g in got] == [
        ("F1", "O1", "E1"), ("F2", "EBSCO", "PHY")]
    used = (tmp_path / "out" / "prep_defaults_used.csv").read_text(encoding="utf-8")
    assert "FOLIO Access Provider,EBSCO" in used and "FOLIO Expense Class,PHY" in used
    assert "FOLIO Fund" not in used
    report = (tmp_path / "out" / "prep_report.txt").read_text(encoding="utf-8")
    assert "FOLIO Access Provider: customer column present; blank on 1 of 2" in report


def test_config_is_split_between_fixed_and_library_settings():
    fixed = json.loads((ROOT / "pipeline" / "pipeline_config.json").read_text(
        encoding="utf-8"))
    library = json.loads((ROOT / "ebsconet_config.json").read_text(encoding="utf-8"))
    overlap = (set(fixed) & set(library)) - {"_comment"}
    assert overlap == set()                         # each setting lives in one file only
    for key in ("columns", "added_columns", "extra_marc_map", "format_routes",
                "output_names", "marc_indicators"):
        assert key in fixed and key not in library
    for key in ("rules", "ongoing", "fund_by_route", "folio", "default_org"):
        assert key in library and key not in fixed


def test_load_config_merges_both_files_and_the_library_wins(tmp_path):
    lib = tmp_path / "lib.json"
    lib.write_text(json.dumps({"rules": {"x": 1}, "output_names": {"online": "mine.xlsx"}}),
                   encoding="utf-8")
    cfg = prep.load_config(lib)
    assert cfg["columns"]["title"] == "Title Name"            # from the fixed file
    assert cfg["rules"] == {"x": 1}                           # from the library file
    assert cfg["output_names"]["online"] == "mine.xlsx"       # library overrides...
    assert cfg["output_names"]["print"] == "library-EBSCONET-print.xlsx"   # ...per key


def test_check_headers_names_missing_columns(cfg):
    prep.check_headers(HEADERS, cfg)
    with pytest.raises(ValueError, match=re.escape("'Total Cost' (cost)")):
        prep.check_headers([h for h in HEADERS if h != "Total Cost"], cfg)


def test_prepare_rejects_sop_with_renamed_heading(cfg, tmp_path):
    src = tmp_path / "sop.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Title", "ISSN", "Format", "Order Number", "Total Cost"])
    wb.save(src)
    with pytest.raises(ValueError, match="'Title Name'"):
        prep.prepare(src, tmp_path / "out", cfg)
    with pytest.raises(ValueError, match="'Title Name'"):
        prep.prepare_for_customer(src, tmp_path / "out2", cfg)
