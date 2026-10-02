import importlib.util
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("ebsconet_configure",
                                              ROOT / "bin" / "ebsconet_configure.py")
cfgmod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cfgmod)

DATA = {
    "/finance/funds": [{"code": "ELEC", "name": "Electronic"},
                       {"code": "PRINT", "name": "Print"}],
    "/finance/expense-classes": [{"code": "SER", "name": "Serials"}],
    "/organizations/organizations": [{"code": "EBSCO", "name": "EBSCO", "isVendor": True},
                                     {"code": "NOTV", "name": "Not vendor", "isVendor": False}],
    "/locations": [{"code": "MAIN", "name": "Main"}, {"code": "ANX", "name": "Annex"}],
    "/material-types": [{"name": "book"}, {"name": "journal"}],
    "/orders/acquisition-methods": [{"value": "Subscription"}],
}


class FakeClient:
    def folio_get(self, path, key=None, query_params=None):
        return DATA[path]


def scripted(answers):
    it = iter(answers)
    return lambda prompt: next(it)


def test_pick_many_and_one():
    items = ["a", "b", "c"]
    assert cfgmod.pick_many(items, str, "x", read=scripted(["1,3"])) == ["a", "c"]
    assert cfgmod.pick_many(items, str, "x", read=scripted([""])) == []
    assert cfgmod.pick_one(items, str, "x", read=scripted(["9", "2"])) == "b"
    assert cfgmod.pick_one(items, str, "x", allow_blank=True, read=scripted([""])) is None


def test_filter_narrows_long_list():
    items = ["item%02d" % n for n in range(40)]
    got = cfgmod.pick_one(items, str, "x", read=scripted(["item3", "1"]))
    assert got == "item30"


def test_main_writes_config(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    base = json.loads((ROOT / "ebsconet_config.json").read_text())
    base["custom_key"] = {"Wiley": "WILEY"}
    base["expense_class_by_subject"] = {"Old": "OLD"}
    (tmp_path / "ebsconet_config.json").write_text(json.dumps(base))
    (tmp_path / "t.ini").write_text("okapiUrl=x\n")
    answers = [
        "1", "2", "2",                  # funds online/print/pe
        "y", "1", "n",                  # use classes, default class, no subject map
        "1",                            # EBSCOnet organization (NOTV filtered out)
        "1,2", "2",                     # location choices, default location
        "2", "1",                       # material types: journal, then default
        "1",                            # acquisition method
        "4",                            # payment method (4 = Deposit Account)
        "One-Time", "180",              # ongoing
        "y",                            # write
    ]
    rc = cfgmod.main(["--ini", "t.ini"], read=scripted(answers),
                     connect_fn=lambda ini: FakeClient())
    assert rc == 0
    out = json.loads((tmp_path / "ebsconet_config.json").read_text())
    assert out["fund_by_route"] == {"online": "ELEC", "print": "PRINT", "pe": "PRINT"}
    assert out["default_expense_class"] == "SER"
    assert out["expense_class_by_subject"] == {}
    assert out["folio"]["vendor_org_code"] == "EBSCO"
    assert out["customer_choices"]["location"] == ["Annex (ANX)", "Main (MAIN)"]
    assert out["folio"]["location"] == "Main (MAIN)"
    assert out["customer_choices"]["material_type"] == ["journal"]
    assert out["folio"]["physical_material_type"] == "journal"
    assert out["folio"]["acquisition_method"] == "Subscription"
    assert out["folio"]["account_payment_method"] == "Deposit Account"
    assert out["ongoing"] == {**base["ongoing"], "default_order_type": "One-Time",
                              "interval_days": 180}
    assert out["custom_key"] == {"Wiley": "WILEY"}
    assert list(tmp_path.glob("ebsconet_config.json.bak-*"))


def test_decline_writes_nothing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "t.ini").write_text("okapiUrl=x\n")
    answers = ["", "", "", "n", "1", "", "", "", "", "", "", "", "", "n"]
    # blank funds keep current; "n" to classes; the EBSCOnet org is required ("1");
    # blanks skip the rest; decline write
    rc = cfgmod.main(["--ini", "t.ini"], read=scripted(answers),
                     connect_fn=lambda ini: FakeClient())
    assert rc == 1
    assert not (tmp_path / "ebsconet_config.json").exists()


def test_worksheet(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.chdir(tmp_path)
    (tmp_path / "t.ini").write_text("okapiUrl=x\n")
    rc = cfgmod.main(["--ini", "t.ini", "--worksheet"], connect_fn=lambda ini: FakeClient())
    assert rc == 0
    wb = load_workbook(tmp_path / cfgmod.WORKSHEET_NAME)
    ws = wb["Defaults"]
    assert wb.sheetnames[0] == "Defaults"
    assert [c.value for c in ws[1]] == ["Setting", "What it is", "Current default", "Your answer"]
    assert ws.max_row == 14
    assert ws["A2"].value == "Default EBSCOnet organization"
    assert [r[0] for r in wb["--Funds"].iter_rows(values_only=True)] == [
        "Funds", "ELEC - Electronic", "PRINT - Print"]
    assert [r[0] for r in wb["--Organizations"].iter_rows(values_only=True)] == [
        "Organizations", "EBSCO - EBSCO"]
    assert len(ws.data_validations.dataValidation) > 5
