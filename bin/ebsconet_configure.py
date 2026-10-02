"""Interactive helper that builds ebsconet_config.json from the codes on the live tenant.

Run it (as `ebsconet-configure`) from the work/ folder. It connects with a tenant .ini,
lists the funds, expense classes, vendor organizations, locations, material types and
acquisition methods that really exist, and asks you to pick the ones this library uses.
Read-only against FOLIO. Your existing ebsconet_config.json is backed up before the new
one is written; keys the wizard does not ask about (for example org_by_publisher) are kept.
"""
import argparse
import copy
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pipeline.folio_common import connect  # noqa: E402

CONFIG_NAME = "ebsconet_config.json"
WORKSHEET_NAME = "ebsconet_config_worksheet.xlsx"
PAYMENT_METHODS = ["Cash", "Credit Card", "EFT", "Deposit Account", "Physical Check",
                   "Bank Draft", "Lockbox", "Other"]
LIST_LIMIT = 25


def ask(prompt, default=None, read=input):
    """Prompt for free text; a blank answer returns `default`."""
    suffix = " [%s]" % default if default not in (None, "") else ""
    answer = read("%s%s: " % (prompt, suffix)).strip()
    return answer if answer else default


def ask_yes_no(prompt, default=True, read=input):
    while True:
        answer = (ask(prompt + " (y/n)", "y" if default else "n", read) or "").lower()
        if answer in ("y", "yes"):
            return True
        if answer in ("n", "no"):
            return False
        print("  Please answer y or n.")


def filtered(items, label, read=input):
    """Return items, narrowed by a text filter while the list is too long to read."""
    while len(items) > LIST_LIMIT:
        text = read("  %d entries. Type part of a name/code to narrow the list "
                    "(blank = show first %d): " % (len(items), LIST_LIMIT)).strip().lower()
        if not text:
            return items[:LIST_LIMIT]
        narrowed = [i for i in items if text in label(i).lower()]
        if not narrowed:
            print("  Nothing matches '%s'." % text)
            continue
        items = narrowed
    return items


def pick_one(items, label, prompt, default=None, allow_blank=False, read=input):
    """Choose one entry by number; returns the item, or None when blank is allowed."""
    if not items:
        print("  (none found on this tenant)")
        return None
    shown = filtered(items, label, read)
    for n, item in enumerate(shown, 1):
        print("  %2d. %s" % (n, label(item)))
    while True:
        extra = ", blank to skip" if allow_blank else ""
        answer = ask("%s (number%s)" % (prompt, extra), default, read)
        if not answer:
            if allow_blank:
                return None
            print("  A choice is required.")
            continue
        if answer.isdigit() and 1 <= int(answer) <= len(shown):
            return shown[int(answer) - 1]
        for item in shown:
            if answer.lower() == label(item).lower():
                return item
        print("  Enter a number from the list.")


def pick_many(items, label, prompt, read=input):
    """Choose several entries ('1,3,5'); blank means none (free-text column)."""
    if not items:
        print("  (none found on this tenant)")
        return []
    shown = filtered(items, label, read)
    for n, item in enumerate(shown, 1):
        print("  %2d. %s" % (n, label(item)))
    while True:
        answer = ask("%s (numbers separated by commas, blank = none)" % prompt, None, read)
        if not answer:
            return []
        try:
            numbers = [int(p) for p in answer.replace(" ", "").split(",") if p]
        except ValueError:
            numbers = []
        if numbers and all(1 <= n <= len(shown) for n in numbers):
            return [shown[n - 1] for n in dict.fromkeys(numbers)]
        print("  Enter numbers from the list, e.g. 1,3.")


def fetch(client, path, key):
    """List a FOLIO reference collection (up to 1000 rows); [] if it cannot be read."""
    try:
        return client.folio_get(path, key=key, query_params={"limit": 1000}) or []
    except Exception as exc:  # network / permission problems should not end the wizard
        print("  Could not read %s: %s" % (path, exc))
        return []


def code_name(item):
    return "%s - %s" % (item.get("code", ""), item.get("name", ""))


def location_text(item):
    return "%s (%s)" % (item["name"], item["code"])


def deep_merge(base, extra):
    for key, value in extra.items():
        if isinstance(value, dict) and isinstance(base.get(key), dict):
            deep_merge(base[key], value)
        else:
            base[key] = value
    return base


def heading(text):
    print("\n== %s ==" % text)


def gather(client, cfg, read=input):
    """Walk through each section; returns the dict of values to merge into cfg."""
    out = {"fund_by_route": {}, "folio": {}, "rules": {}, "ongoing": {},
           "customer_choices": {}}
    current = cfg

    heading("1/8 Funds (default fund when the customer leaves a cell blank)")
    funds = sorted(fetch(client, "/finance/funds", "funds"), key=lambda f: f.get("code", ""))
    for route, what in (("online", "electronic"), ("print", "print"), ("pe", "print+electronic")):
        old = current.get("fund_by_route", {}).get(route)
        print("Fund for %s lines (current: %s)" % (what, old))
        fund = pick_one(funds, code_name, "Fund", allow_blank=True, read=read)
        out["fund_by_route"][route] = fund["code"] if fund else old
    print("  Note: each fund needs an Active budget for the current fiscal year.")

    heading("2/8 Expense classes")
    classes = sorted(fetch(client, "/finance/expense-classes", "expenseClasses"),
                     key=lambda c: c.get("code", ""))
    use = ask_yes_no("Does this library use expense classes on its orders?",
                     current.get("rules", {}).get("use_expense_classes", True), read)
    out["rules"]["use_expense_classes"] = use
    if use:
        default = pick_one(classes, code_name, "Default expense class", allow_blank=True,
                           read=read)
        out["default_expense_class"] = (default["code"] if default
                                        else current.get("default_expense_class"))
        mapping = {}
        if ask_yes_no("Map SOP subjects to expense classes (e.g. 'Physical Sciences')?",
                      False, read):
            while True:
                subject = ask("SOP subject text (blank to finish)", None, read)
                if not subject:
                    break
                cls = pick_one(classes, code_name, "Expense class for '%s'" % subject,
                               allow_blank=True, read=read)
                if cls:
                    mapping[subject] = cls["code"]
        out["expense_class_by_subject"] = mapping
    else:
        out["default_expense_class"] = ""
        out["expense_class_by_subject"] = {}

    heading("3/8 Vendor organization")
    orgs = [o for o in fetch(client, "/organizations/organizations", "organizations")
            if o.get("isVendor", True)]
    orgs.sort(key=lambda o: o.get("code", ""))
    print("Default vendor organization when the customer leaves 'FOLIO Org' blank "
          "(current: %s)" % current.get("default_org"))
    org = pick_one(orgs, code_name, "Organization", allow_blank=True, read=read)
    if org:
        out["default_org"] = org["code"]
    print("Organization that vendor accounts are added to by folio_setup "
          "(current: %s)" % current.get("folio", {}).get("vendor_org_code"))
    vendor = pick_one(orgs, code_name, "Organization", allow_blank=True, read=read)
    if vendor:
        out["folio"]["vendor_org_code"] = vendor["code"]

    heading("4/8 Locations (physical and P-E lines)")
    locations = sorted(fetch(client, "/locations", "locations"), key=lambda x: x.get("name", ""))
    print("Locations the customer may choose from (drop-down); blank = free text:")
    chosen = pick_many(locations, location_text, "Locations", read=read)
    out["customer_choices"]["location"] = [location_text(x) for x in chosen]
    print("Default location when the customer leaves it blank:")
    default_loc = pick_one(chosen or locations, location_text, "Default location",
                           allow_blank=True, read=read)
    if default_loc:
        out["folio"]["location"] = location_text(default_loc)

    heading("5/8 Material types (physical and P-E lines)")
    mtypes = sorted(fetch(client, "/material-types", "mtypes"), key=lambda x: x.get("name", ""))
    name = lambda m: m["name"]  # noqa: E731
    print("Material types the customer may choose from (drop-down); blank = free text:")
    chosen_mt = pick_many(mtypes, name, "Material types", read=read)
    out["customer_choices"]["material_type"] = [m["name"] for m in chosen_mt]
    print("Default material type when the customer leaves it blank:")
    default_mt = pick_one(chosen_mt or mtypes, name, "Default material type",
                          allow_blank=True, read=read)
    if default_mt:
        out["folio"]["physical_material_type"] = default_mt["name"]

    heading("6/8 Acquisition method")
    methods = sorted(fetch(client, "/orders/acquisition-methods", "acquisitionMethods"),
                     key=lambda x: x.get("value", ""))
    print("Acquisition method put on every order line (current: %s)"
          % current.get("folio", {}).get("acquisition_method"))
    method = pick_one(methods, lambda m: m["value"], "Acquisition method",
                      allow_blank=True, read=read)
    if method:
        out["folio"]["acquisition_method"] = method["value"]

    heading("7/8 Vendor account payment method")
    print("Payment method for vendor accounts added from the SOP's Account Number "
          "(current: %s)" % current.get("folio", {}).get("account_payment_method"))
    pay = pick_one(PAYMENT_METHODS, str, "Payment method", allow_blank=True, read=read)
    if pay:
        out["folio"]["account_payment_method"] = pay

    heading("8/8 Ongoing order defaults")
    order_type = ask("Default order type when the customer leaves it blank (Ongoing/One-Time)",
                     current.get("ongoing", {}).get("default_order_type", "Ongoing"), read)
    out["ongoing"]["default_order_type"] = (
        "One-Time" if str(order_type).lower().replace("-", "").replace(" ", "") == "onetime"
        else "Ongoing")
    while True:
        days = ask("Default renewal interval in days",
                   str(current.get("ongoing", {}).get("interval_days", 365)), read)
        if str(days).isdigit() and int(days) > 0:
            out["ongoing"]["interval_days"] = int(days)
            break
        print("  Enter a whole number of days.")
    return out


def write_worksheet(client, cfg, path):
    """Write an .xlsx the customer can fill in: one row per setting, a drop-down of the
    tenant's real options where the answer is a single choice, and the lists on extra
    sheets. The yellow column is theirs to fill."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    def srt(rows, field):
        return sorted(rows, key=lambda r: r.get(field, ""))

    funds = [code_name(f) for f in srt(fetch(client, "/finance/funds", "funds"), "code")]
    classes = [code_name(c) for c in srt(
        fetch(client, "/finance/expense-classes", "expenseClasses"), "code")]
    orgs = [code_name(o) for o in srt(
        [o for o in fetch(client, "/organizations/organizations", "organizations")
         if o.get("isVendor", True)], "code")]
    locations = [location_text(x) for x in srt(fetch(client, "/locations", "locations"), "name")]
    mtypes = [m["name"] for m in srt(fetch(client, "/material-types", "mtypes"), "name")]
    methods = [m["value"] for m in srt(
        fetch(client, "/orders/acquisition-methods", "acquisitionMethods"), "value")]
    lists = {"Funds": funds, "ExpenseClasses": classes, "Organizations": orgs,
             "Locations": locations, "MaterialTypes": mtypes,
             "AcquisitionMethods": methods, "PaymentMethods": PAYMENT_METHODS,
             "OrderTypes": ["Ongoing", "One-Time"], "YesNo": ["Yes", "No"]}

    cur = cfg
    fund = cur.get("fund_by_route", {})
    folio = cur.get("folio", {})
    # (setting, what it is, list sheet for a drop-down or None, kind, current value)
    rows = [
        ("Fund: electronic", "Fund charged for electronic (online) subscriptions when the "
         "line has no fund of its own. Needs an Active budget this fiscal year.",
         "Funds", "one", fund.get("online")),
        ("Fund: print", "Fund charged for print subscriptions.", "Funds", "one",
         fund.get("print")),
        ("Fund: print + electronic", "Fund charged for combined print + electronic "
         "subscriptions.", "Funds", "one", fund.get("pe")),
        ("Use expense classes?", "Does the library put an expense class on its orders?",
         "YesNo", "one", "Yes" if cur.get("rules", {}).get("use_expense_classes", True)
         else "No"),
        ("Default expense class", "Expense class used when none is given (skip if the "
         "answer above is No).", "ExpenseClasses", "one", cur.get("default_expense_class")),
        ("Subject to expense class (optional)", "Optional: one line per SOP subject, written "
         "Subject = class code, e.g. Physical Sciences = SER.", None, "text", None),
        ("Default vendor organization", "Organization used for a subscription when no "
         "organization is given.", "Organizations", "one", cur.get("default_org")),
        ("Organization for vendor accounts", "Organization the SOP account numbers are "
         "added to.", "Organizations", "one", folio.get("vendor_org_code")),
        ("Locations offered to the customer", "Locations that appear as a drop-down on the "
         "physical and print+electronic spreadsheets. List every one wanted, separated by "
         "semicolons (see the Locations sheet). Blank = free text.", None, "text",
         "; ".join(cur.get("customer_choices", {}).get("location", [])) or None),
        ("Default location", "Used when the customer leaves the location blank.",
         "Locations", "one", folio.get("location")),
        ("Material types offered to the customer", "Material types for the drop-down on the "
         "physical and print+electronic spreadsheets, separated by semicolons (see the "
         "MaterialTypes sheet). Blank = free text.", None, "text",
         "; ".join(cur.get("customer_choices", {}).get("material_type", [])) or None),
        ("Default material type", "Used when the customer leaves it blank.",
         "MaterialTypes", "one", folio.get("physical_material_type")),
        ("Acquisition method", "Acquisition method put on every order line.",
         "AcquisitionMethods", "one", folio.get("acquisition_method")),
        ("Vendor account payment method", "Payment method for vendor accounts added from "
         "the SOP's Account Number column.", "PaymentMethods", "one",
         folio.get("account_payment_method")),
        ("Default order type", "Order type when the customer leaves it blank.",
         "OrderTypes", "one", cur.get("ongoing", {}).get("default_order_type")),
        ("Default renewal interval (days)", "Days between renewals of an ongoing order, "
         "when left blank (365 = yearly).", None, "text",
         cur.get("ongoing", {}).get("interval_days")),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Worksheet"
    ws.append(["Setting", "What it is", "Your answer"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFFFCC")
    list_ws = {}
    for name, items in lists.items():
        sheet = wb.create_sheet(name)
        sheet.append([name])
        sheet["A1"].font = Font(bold=True)
        for item in items:
            sheet.append([item])
        sheet.column_dimensions["A"].width = 60
        list_ws[name] = len(items)
    for n, (setting, text, source, kind, value) in enumerate(rows, 2):
        ws.append([setting, text, ""])
        ws.cell(n, 3).fill = yellow
        if source and list_ws[source]:
            dv = DataValidation(type="list", allow_blank=True, showErrorMessage=False,
                                formula1="=%s!$A$2:$A$%d" % (source, list_ws[source] + 1))
            ws.add_data_validation(dv)
            dv.add(ws.cell(n, 3))
    for col, width in zip("ABC", (38, 70, 50)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def load_config(path, fallback):
    for candidate in (path, fallback):
        if candidate and Path(candidate).exists():
            return json.loads(Path(candidate).read_text(encoding="utf-8")), candidate
    raise SystemExit("No %s found in %s or %s." % (CONFIG_NAME, path.parent, fallback))


def choose_ini(read=input):
    found = sorted(p.name for p in Path.cwd().glob("*.ini") if p.name != "sample.ini")
    if len(found) == 1:
        print("Using tenant file %s" % found[0])
        return found[0]
    if found:
        print("Tenant files here: " + ", ".join(found))
    else:
        print("No <tenant>.ini here. Copy sample.ini to <tenant>.ini and fill it in first.")
    return ask("Tenant .ini file", found[0] if found else None, read)


def main(argv=None, read=input, connect_fn=connect):
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--ini", help="tenant .ini file (default: ask)")
    parser.add_argument("--config", default=CONFIG_NAME,
                        help="config file to create/update (default: %(default)s)")
    parser.add_argument("--worksheet", nargs="?", const=WORKSHEET_NAME, metavar="FILE",
                        help="write an Excel worksheet of every question and the tenant's "
                             "options for the customer to fill in, instead of asking "
                             "(default file: %s)" % WORKSHEET_NAME)
    args = parser.parse_args(argv)

    app_default = Path(__file__).resolve().parent.parent / CONFIG_NAME
    config_path = Path(args.config)
    cfg, source = load_config(config_path, app_default)
    print("Starting from %s" % source)

    ini = args.ini or choose_ini(read)
    if not ini or not Path(ini).exists():
        raise SystemExit("Tenant file %r not found." % ini)
    print("Connecting...")
    client = connect_fn(ini)
    if args.worksheet:
        path = write_worksheet(client, cfg, args.worksheet)
        print("Wrote worksheet %s. Send it to the customer; when it comes back, run "
              "ebsconet-configure and enter their answers." % path)
        return 0
    print("Connected. Every question below lists what exists on this tenant.\n"
          "Press Enter to keep the current value where one is shown.")

    answers = gather(client, cfg, read)
    new = deep_merge(copy.deepcopy(cfg), answers)
    new["expense_class_by_subject"] = answers["expense_class_by_subject"]  # replace, not merge
    print("\n== Result ==")
    print(json.dumps(new, indent=2))
    if not ask_yes_no("\nWrite %s?" % config_path, True, read):
        print("Nothing written.")
        return 1
    if config_path.exists():
        backup = config_path.with_name("%s.bak-%s" % (
            config_path.name, datetime.now().strftime("%Y%m%d_%H%M%S")))
        shutil.copy2(config_path, backup)
        print("Backed up the old file to %s" % backup)
    config_path.write_text(json.dumps(new, indent=2) + "\n", encoding="utf-8")
    print("Wrote %s. Next: ebsconet for-customer SOP.xlsx (customer drop-downs now use "
          "your locations and material types)." % config_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
