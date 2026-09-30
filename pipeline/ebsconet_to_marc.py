"""Convert the prepared EBSCONET workbooks into MARC files for FOLIO Data Import.

Replaces the MarcEdit delimited-text step (instruction steps 13-22). The SOP column ->
MARC tag map is read from order_marc_headers.xlsx (row 1 = tag$code or "Ignore",
row 2 = SOP column name); columns added by ebsconet_prep.py are mapped from the
"extra_marc_map" section of the config. Writes a .mrc (for loading) and a .mrk
(MarcEdit-style text, for inspection) per workbook.
"""
import argparse
import io
from pathlib import Path

from openpyxl import load_workbook
from pymarc import Field, MARCReader, Record, Subfield

from pipeline.ebsconet_prep import blank, load_config

BLANK = (" ", " ")


def parse_target(target):
    """'990$s' -> ('990', 's'); returns None for Ignore/blank/malformed targets."""
    if target is None:
        return None
    text = str(target).strip()
    if not text or text.lower() == "ignore" or "$" not in text:
        return None
    tag, code = text.split("$", 1)
    if len(tag) != 3 or len(code) != 1:
        return None
    return tag, code


def load_tag_map(headers_file, cfg):
    """Return ordered list of (sop_column, tag, code)."""
    wb = load_workbook(headers_file, read_only=True, data_only=True)
    tags, names = wb.worksheets[0].iter_rows(min_row=1, max_row=2, values_only=True)
    wb.close()
    tag_map = []
    for target, name in zip(tags, names):
        parsed = parse_target(target)
        if parsed and name:
            tag_map.append((str(name).strip(),) + parsed)
    for name, target in cfg["extra_marc_map"].items():
        parsed = parse_target(target)
        if parsed:
            tag_map.append((name,) + parsed)
    return tag_map


def fmt_value(value):
    if blank(value):
        return ""
    if isinstance(value, float):
        return "%.2f" % value
    return str(value).strip()


def build_record(row, tag_map, control_id, cfg):
    """Build one pymarc Record from a workbook row (dict of column -> value)."""
    grouped = {}
    for column, tag, code in tag_map:
        value = fmt_value(row.get(column))
        if value:
            grouped.setdefault(tag, []).append(Subfield(code=code, value=value))
    rec = Record(force_utf8=True)
    rec.leader = "00000nas a2200000   4500"
    rec.add_field(Field(tag="001", data=control_id))
    for tag in sorted(grouped):
        ind = cfg["marc_indicators"].get(tag, BLANK)
        rec.add_field(Field(tag=tag, indicators=list(ind), subfields=grouped[tag]))
    return rec


def read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = [dict(zip(headers, v)) for v in it
            if not all(x is None or x == "" for x in v)]
    wb.close()
    return rows


def control_ids(rows, order_column):
    """001 values: order number, suffixed -2, -3... for repeated order numbers."""
    seen = {}
    ids = []
    for i, row in enumerate(rows, start=1):
        order = fmt_value(row.get(order_column)) or "ebsconet-row%d" % i
        seen[order] = seen.get(order, 0) + 1
        ids.append(order if seen[order] == 1 else "%s-%d" % (order, seen[order]))
    return ids


def _mrk_escape(text):
    return text.replace("$", "{dollar}")


def to_mrk(rec):
    lines = ["=LDR  " + str(rec.leader)]
    for f in rec.get_fields():
        if f.is_control_field():
            lines.append("=%s  %s" % (f.tag, f.data))
        else:
            ind = "".join("\\" if i == " " else i for i in f.indicators)
            subs = "".join("$%s%s" % (s.code, _mrk_escape(s.value))
                           for s in f.subfields)
            lines.append("=%s  %s%s" % (f.tag, ind, subs))
    return "\n".join(lines) + "\n"


def verify_round_trip(mrc_path, expected):
    """Re-read the .mrc and compare with the records we meant to write."""
    problems = []
    with open(mrc_path, "rb") as fh:
        got = list(MARCReader(fh, to_unicode=True, force_utf8=True))
    if len(got) != len(expected):
        return ["record count %d != %d" % (len(got), len(expected))]
    for n, (g, e) in enumerate(zip(got, expected), start=1):
        # the leader is recalculated on write, so compare everything after it
        if to_mrk(g).splitlines()[1:] != to_mrk(e).splitlines()[1:]:
            problems.append("record %d differs after round trip" % n)
    return problems


def convert_workbook(xlsx_path, out_prefix, tag_map, cfg):
    """Write out_prefix.mrc / .mrk for one workbook. Returns a summary dict."""
    rows = read_rows(xlsx_path)
    if not rows:
        for ext in (".mrc", ".mrk"):
            Path(out_prefix).with_suffix(ext).unlink(missing_ok=True)
        return {"file": Path(out_prefix).with_suffix(".mrc").name, "records": 0,
                "warnings": [], "problems": []}
    ids = control_ids(rows, cfg["columns"]["order_number"])
    records = [build_record(r, tag_map, cid, cfg) for r, cid in zip(rows, ids)]
    warnings = ["row %d: no 245$a title" % i
                for i, r in enumerate(records, start=2) if r.get("245") is None]
    out_prefix = Path(out_prefix)
    out_prefix.parent.mkdir(parents=True, exist_ok=True)
    mrc = out_prefix.with_suffix(".mrc")
    with open(mrc, "wb") as fh:
        for rec in records:
            fh.write(rec.as_marc())
    with io.open(out_prefix.with_suffix(".mrk"), "w", encoding="utf-8",
                 newline="\n") as fh:
        fh.write("\n".join(to_mrk(r) for r in records))
    problems = verify_round_trip(mrc, records)
    return {"file": mrc.name, "records": len(records), "warnings": warnings,
            "problems": problems}


def convert_all(in_dir, cfg, headers_file, out_dir=None):
    in_dir = Path(in_dir)
    out_dir = Path(out_dir or in_dir / cfg["marc_output_dir"])
    tag_map = load_tag_map(headers_file, cfg)
    results = []
    for name in cfg["output_names"].values():
        src = in_dir / name
        if src.exists():
            results.append(convert_workbook(src, out_dir / Path(name).stem,
                                            tag_map, cfg))
    return results


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--in-dir", default="out", help="folder holding the prep workbooks")
    p.add_argument("--config", default="ebsconet_config.json")
    p.add_argument("--headers", help="order_marc_headers.xlsx (default from config)")
    p.add_argument("--out", help="output folder (default <in-dir>/marc)")
    args = p.parse_args(argv)
    cfg = load_config(args.config)
    results = convert_all(args.in_dir, cfg, args.headers or cfg["headers_file"],
                          args.out)
    bad = False
    for r in results:
        print("%s: %d records" % (r["file"], r["records"]))
        for w in r["warnings"]:
            print("  WARNING", w)
        for pr in r["problems"]:
            print("  ERROR", pr)
            bad = True
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
