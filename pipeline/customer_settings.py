"""The "Defaults" sheet: the library's answers to the setup questions, in one workbook.

`ebsconet-configure --worksheet` writes it on its own; `ebsconet for-customer` writes it as
the first sheet of the customer workbook, ahead of the electronic / physical / P-E sheets.
Both ask the same questions (`settings_rows`) and list the tenant's real options on extra
sheets (`LIST_SHEETS`). `read_overrides` turns the filled-in sheet back into the config
keys it answers, so `ebsconet build` can use them.
"""
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SETTINGS_SHEET = "Defaults"
PAYMENT_METHODS = ["Cash", "Credit Card", "EFT", "Deposit Account", "Physical Check",
                   "Bank Draft", "Lockbox", "Other"]
LIST_SHEETS = ("Funds", "ExpenseClasses", "Organizations", "Locations", "MaterialTypes",
               "AcquisitionMethods", "PaymentMethods", "OrderTypes", "YesNo")
LIST_PREFIX = "--"      # marks sheets that only feed drop-downs, not sheets to fill in
YELLOW = PatternFill("solid", fgColor="FFFFCC")


def fetch(client, path, key):
    """List a FOLIO reference collection (up to 1000 rows); [] if it cannot be read."""
    try:
        return client.folio_get(path, key=key, query_params={"limit": 1000}) or []
    except Exception as exc:  # network / permission problems should not stop the caller
        print("  Could not read %s: %s" % (path, exc))
        return []


def code_name(item):
    return "%s - %s" % (item.get("code", ""), item.get("name", ""))


def location_text(item):
    return "%s (%s)" % (item["name"], item["code"])


def _sorted(rows, field):
    return sorted(rows, key=lambda r: r.get(field, ""))


def tenant_lists(client=None):
    """{list sheet name: [options]}. Without a client only the fixed lists are filled."""
    lists = {name: [] for name in LIST_SHEETS}
    lists["PaymentMethods"] = list(PAYMENT_METHODS)
    lists["OrderTypes"] = ["Ongoing", "One-Time"]
    lists["YesNo"] = ["Yes", "No"]
    if client is None:
        return lists
    lists["Funds"] = [code_name(f) for f in _sorted(
        fetch(client, "/finance/funds", "funds"), "code")]
    lists["ExpenseClasses"] = [code_name(c) for c in _sorted(
        fetch(client, "/finance/expense-classes", "expenseClasses"), "code")]
    lists["Organizations"] = [code_name(o) for o in _sorted(
        [o for o in fetch(client, "/organizations/organizations", "organizations")
         if o.get("isVendor", True)], "code")]
    lists["Locations"] = [location_text(x) for x in _sorted(
        fetch(client, "/locations", "locations"), "name")]
    lists["MaterialTypes"] = [m["name"] for m in _sorted(
        fetch(client, "/material-types", "mtypes"), "name")]
    lists["AcquisitionMethods"] = [m["value"] for m in _sorted(
        fetch(client, "/orders/acquisition-methods", "acquisitionMethods"), "value")]
    return lists


def list_range(name, lists):
    """Formula for a drop-down over list sheet `name`, or None if that list is empty."""
    count = len(lists.get(name) or [])
    return "='%s%s'!$A$2:$A$%d" % (LIST_PREFIX, name, count + 1) if count else None


def settings_rows(cfg):
    """[(setting, what it is, list sheet for a drop-down or None, current value)]."""
    fund = cfg.get("fund_by_route", {})
    folio = cfg.get("folio", {})
    return [
        ("Fund: electronic", "Fund charged for electronic (online) subscriptions when the "
         "line has no fund of its own. Needs an Active budget this fiscal year.",
         "Funds", fund.get("online")),
        ("Fund: print", "Fund charged for print subscriptions.", "Funds", fund.get("print")),
        ("Fund: print + electronic", "Fund charged for combined print + electronic "
         "subscriptions.", "Funds", fund.get("pe")),
        ("Use expense classes?", "Does the library put an expense class on its orders?",
         "YesNo", "Yes" if cfg.get("rules", {}).get("use_expense_classes", True) else "No"),
        ("Default expense class", "Expense class used when none is given (skip if the "
         "answer above is No).", "ExpenseClasses", cfg.get("default_expense_class")),
        ("Subject to expense class (optional)", "Optional: one line per SOP subject, written "
         "Subject = class code, e.g. Physical Sciences = SER.", None, None),
        ("Default vendor organization", "Organization used for a subscription when no "
         "organization is given.", "Organizations", cfg.get("default_org")),
        ("Organization for vendor accounts", "Organization the SOP account numbers are "
         "added to.", "Organizations", folio.get("vendor_org_code")),
        ("Default location", "Used when the customer leaves the location blank on the "
         "physical and print + electronic sheets.", "Locations", folio.get("location")),
        ("Default material type", "Used when the customer leaves it blank on the physical "
         "and print + electronic sheets.", "MaterialTypes", folio.get("physical_material_type")),
        ("Acquisition method", "Acquisition method put on every order line.",
         "AcquisitionMethods", folio.get("acquisition_method")),
        ("Vendor account payment method", "Payment method for vendor accounts added from "
         "the SOP's Account Number column.", "PaymentMethods",
         folio.get("account_payment_method")),
        ("Default order type", "Order type when the customer leaves it blank and the SOP "
         "has no Term column (otherwise a line with a Term is Ongoing, one without is "
         "One-Time).",
         "OrderTypes", cfg.get("ongoing", {}).get("default_order_type")),
        ("Default renewal interval (days)", "Days between renewals of an ongoing order, "
         "when left blank (365 = yearly).", None, cfg.get("ongoing", {}).get("interval_days")),
    ]


def add_settings_sheets(wb, cfg, lists, first=True):
    """Add the Defaults sheet (first in the workbook unless first=False) and one sheet per
    non-empty option list (last). Returns the Defaults sheet."""
    ws = wb.create_sheet(SETTINGS_SHEET, 0 if first else None)
    ws.append(["Setting", "What it is", "Current default", "Your answer"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for name in LIST_SHEETS:
        if lists.get(name):
            sheet = wb.create_sheet(LIST_PREFIX + name)
            sheet.append([name])
            sheet["A1"].font = Font(bold=True)
            for item in lists[name]:
                sheet.append([item])
            sheet.column_dimensions["A"].width = 60
    for n, (setting, text, source, current) in enumerate(settings_rows(cfg), 2):
        ws.append([setting, text, "" if current is None else str(current), ""])
        ws.cell(n, 4).fill = YELLOW
        formula = list_range(source, lists) if source else None
        if formula:
            dv = DataValidation(type="list", allow_blank=True, showErrorMessage=False,
                                formula1=formula)
            ws.add_data_validation(dv)
            dv.add(ws.cell(n, 4))
    for col, width in zip("ABCD", (38, 70, 30, 50)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def _code(text):
    """'CODE - Name' -> 'CODE'."""
    return str(text).split(" - ", 1)[0].strip()


def _subjects(text):
    out = {}
    for part in re.split(r"[\n;]", str(text)):
        if "=" in part:
            subject, code = part.split("=", 1)
            if subject.strip() and code.strip():
                out[subject.strip()] = code.strip()
    return out


def read_overrides(path):
    """Answers on the Defaults sheet as config keys ({} if the file has no such sheet or
    nothing is answered). Blank answers are left out so the config default stays."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SETTINGS_SHEET not in wb.sheetnames:
            return {}
        answers = {}
        for row in wb[SETTINGS_SHEET].iter_rows(min_row=2, values_only=True):
            if row and row[0] and len(row) > 3 and row[3] not in (None, ""):
                answers[str(row[0]).strip()] = str(row[3]).strip()
    finally:
        wb.close()
    out = {}

    def put(path_, value):
        node = out
        for key in path_[:-1]:
            node = node.setdefault(key, {})
        node[path_[-1]] = value

    for key, route in (("Fund: electronic", "online"), ("Fund: print", "print"),
                       ("Fund: print + electronic", "pe")):
        if key in answers:
            put(("fund_by_route", route), _code(answers[key]))
    if "Use expense classes?" in answers:
        put(("rules", "use_expense_classes"), answers["Use expense classes?"].lower() == "yes")
    if "Default expense class" in answers:
        put(("default_expense_class",), _code(answers["Default expense class"]))
    if "Subject to expense class (optional)" in answers:
        put(("expense_class_by_subject",),
            _subjects(answers["Subject to expense class (optional)"]))
    if "Default vendor organization" in answers:
        put(("default_org",), _code(answers["Default vendor organization"]))
    if "Organization for vendor accounts" in answers:
        put(("folio", "vendor_org_code"), _code(answers["Organization for vendor accounts"]))
    for key, name in (("Default location", "location"),
                      ("Default material type", "physical_material_type"),
                      ("Acquisition method", "acquisition_method"),
                      ("Vendor account payment method", "account_payment_method")):
        if key in answers:
            put(("folio", name), answers[key])
    if "Default order type" in answers:
        flat = answers["Default order type"].lower().replace("-", "").replace(" ", "")
        put(("ongoing", "default_order_type"), "One-Time" if flat == "onetime" else "Ongoing")
    days = answers.get("Default renewal interval (days)", "")
    if days.replace(".0", "").isdigit() and int(float(days)) > 0:
        put(("ongoing", "interval_days"), int(float(days)))
    return out
